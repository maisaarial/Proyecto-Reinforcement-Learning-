'''
Curriculum Learning
Paso 1: encontrar objeto
Paso 2: encontrar + tomar
Paso 3: encontrar + tomar + salir
'''

# train_curriculum.py
import os
import time
import numpy as np
import gymnasium as gym
from openpyxl import Workbook, load_workbook

from stable_baselines3 import PPO
from stable_baselines3.common.vec_env import DummyVecEnv, VecMonitor
from stable_baselines3.common.monitor import Monitor
from stable_baselines3.common.callbacks import EvalCallback, CheckpointCallback
from stable_baselines3.common.utils import set_random_seed
from stable_baselines3.common.evaluation import evaluate_policy
from gymnasium.wrappers import TimeLimit, RecordVideo


# Importa los 3 entornos (tus scripts)
from Robbery_v1a import ThiefEnv_V1a
from Robbery_v1b import ThiefEnv_V1b
from Robbery_v1c import ThiefEnv_V1c


# ============================================================
# WRAPPERS (adaptan observaciones para SB3)
# ============================================================

class FlattenGridObs(gym.ObservationWrapper):
    """
    Convierte obs["grid"] de (3,3) a (9,) para MLP.
    """
    def __init__(self, env):
        super().__init__(env)
        assert isinstance(self.observation_space, gym.spaces.Dict)

        old = self.observation_space
        grid_space = old["grid"]
        assert isinstance(grid_space, gym.spaces.Box)

        self.observation_space = gym.spaces.Dict({
            **{k: v for k, v in old.spaces.items() if k != "grid"},
            "grid": gym.spaces.Box(
                low=grid_space.low.min(),
                high=grid_space.high.max(),
                shape=(int(np.prod(grid_space.shape)),),
                dtype=grid_space.dtype
            ),
        })

    def observation(self, obs):
        obs = dict(obs)
        obs["grid"] = obs["grid"].reshape(-1)
        return obs


class CastObsToFloat32(gym.ObservationWrapper):
    """
    Convierte Boxes a float32 (SB3 suele ir mejor con float32).
    """
    def __init__(self, env):
        super().__init__(env)
        assert isinstance(self.observation_space, gym.spaces.Dict)

        new_spaces = {}
        for k, sp in self.observation_space.spaces.items():
            if isinstance(sp, gym.spaces.Box):
                new_spaces[k] = gym.spaces.Box(
                    low=sp.low, high=sp.high, shape=sp.shape, dtype=np.float32
                )
            else:
                new_spaces[k] = sp
        self.observation_space = gym.spaces.Dict(new_spaces)

    def observation(self, obs):
        obs = dict(obs)
        for k, v in obs.items():
            if isinstance(v, np.ndarray) and v.dtype != np.float32:
                obs[k] = v.astype(np.float32)
        return obs


class RemoveKeysObs(gym.ObservationWrapper):
    """
    Elimina claves de la observación (y del observation_space).

    Esto es CLAVE para curriculum si alguna versión (v1c) devuelve una key extra
    (por ejemplo "exits") pero tú quieres mantener el mismo input al policy
    entre todas las etapas.
    """
    def __init__(self, env, keys_to_remove=("exits",)):
        super().__init__(env)
        assert isinstance(self.observation_space, gym.spaces.Dict)

        self.keys_to_remove = set(keys_to_remove)
        old = self.observation_space

        self.observation_space = gym.spaces.Dict({
            k: v for k, v in old.spaces.items() if k not in self.keys_to_remove
        })

    '''
    def observation(self, obs):
        obs = dict(obs)
        for k in self.keys_to_remove:
            obs.pop(k, None)
        return obs
    '''

    def observation(self, obs):
        # Si no es dict, te lo decimos clarito
        if not isinstance(obs, dict):
            raise TypeError(f"RemoveKeysObs recibió obs de tipo {type(obs)} con valor: {repr(obs)[:200]}")

        # ya es dict, podemos copiarlo
        obs = dict(obs)
        for k in self.keys_to_remove:
            obs.pop(k, None)
        return obs

# ============================================================
# FACTORY: crea entornos para vectorización (cambiando env_cls)
# ============================================================

def make_env(env_cls, rank: int, seed: int, render: bool = False):
    """
    env_cls: clase del entorno (ThiefEnv_V1a / V1b / V1c)
    """
    def _init():
        env = env_cls(render_mode="human" if render else None)
        env = TimeLimit(env,max_episode_steps=250)
        # Si el env trae keys extras (p.ej. v1c trae "exits"), las quitamos
        env = RemoveKeysObs(env, keys_to_remove=("exits",))

        # Preprocesado de observaciones
        env = FlattenGridObs(env)
        env = CastObsToFloat32(env)

        # Seed por entorno
        env.reset(seed=seed + rank)
        return env

    set_random_seed(seed)
    return _init

#Guardar excels 
'''
def get_last_eval_mean_std(eval_dir: str):
    """
    Lee el último evaluation guardado por EvalCallback en eval_dir/evaluations.npz
    Devuelve (mean, std). Si no existe, devuelve (None, None).
    """
    npz_path = os.path.join(eval_dir, "evaluations.npz")
    if not os.path.exists(npz_path):
        return None, None

    data = np.load(npz_path)
    # 'results' suele ser shape (n_evals, n_eval_episodes)
    results = data["results"]
    if results.size == 0:
        return None, None

    last = results[-1]  # rewards del último bloque de evaluación
    return float(np.mean(last)), float(np.std(last))
'''

def append_row_to_excel(xlsx_path: str, sheet_name: str, row_dict: dict):
    """
    Crea el excel si no existe. Si existe, añade una fila al final sin borrar nada.
    row_dict: {"col1": val1, "col2": val2, ...}
    """
    if os.path.exists(xlsx_path):
        wb = load_workbook(xlsx_path)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

    # Si está vacío, escribe cabeceras
    if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
        headers = list(row_dict.keys())
        ws.append(headers)

    # Asegura el mismo orden de columnas que el header
    headers = [cell.value for cell in ws[1]]
    row = [row_dict.get(h, "") for h in headers]
    ws.append(row)

    wb.save(xlsx_path)

#Para el video 
def record_one_episode(model, env_cls, video_dir, stage_name, seed=0):
    os.makedirs(video_dir, exist_ok=True)

    # Creamos el env base
    env = env_cls(render_mode="rgb_array")
    
    # Aplicamos wrappers manualmente pero sin repetir el render_mode
    env = TimeLimit(env, max_episode_steps=250)
    env = RemoveKeysObs(env, keys_to_remove=("exits",))
    env = FlattenGridObs(env)
    env = CastObsToFloat32(env)
    
    # El wrapper de video debe ser el último
    env = RecordVideo(
        env, 
        video_folder=video_dir, 
        episode_trigger=lambda e: True,
        name_prefix=f"{stage_name}_{int(time.time())}" # Esto ayuda con el error de 'Overwriting'
    )

    obs, info = env.reset(seed=seed)
    terminated = truncated = False
    
    while not (terminated or truncated):
        action, _ = model.predict(obs, deterministic=True)
        # Asegúrate de pasar la acción como int si es Discrete
        obs, reward, terminated, truncated, info = env.step(int(action))

    env.close()

    '''
def record_one_episode(model, env_cls, video_dir, seed=0):
    os.makedirs(video_dir, exist_ok=True)

    # 1. Crear el entorno con render_mode
    env = env_cls(render_mode="rgb_array")
    
    # 2. APLICAR LOS MISMOS WRAPPERS QUE EN EL ENTRENAMIENTO
    env = TimeLimit(env, max_episode_steps=250)
    env = RemoveKeysObs(env, keys_to_remove=("exits",))
    env = FlattenGridObs(env)
    env = CastObsToFloat32(env)
    
    # 3. Añadir el wrapper de vídeo al final
    env = RecordVideo(env, video_folder=video_dir, episode_trigger=lambda e: True, name_prefix="final_eval")

    obs, info = env.reset(seed=seed)
    terminated = truncated = False
    
    while not (terminated or truncated):
        # Ahora obs tendrá forma (9,) y el modelo no fallará
        action, _ = model.predict(obs, deterministic=True)
        obs, reward, terminated, truncated, info = env.step(int(action))

    env.close()
    
def record_one_episode(model, env_cls, video_dir, seed=0):
    os.makedirs(video_dir, exist_ok=True)

    env = env_cls(render_mode="rgb_array")
    env = TimeLimit(env, max_episode_steps=250)
    env = RecordVideo(env, video_folder=video_dir, episode_trigger=lambda e: True)

    obs, info = env.reset(seed=seed)
    terminated = truncated = False
    while not (terminated or truncated):
        action, _ = model.predict(obs, deterministic=True)
        obs, reward, terminated, truncated, info = env.step(int(action))

    env.close()
'''

# ============================================================
# MAIN (curriculum por etapas)
# ============================================================

def main():
    # -------------------------
    # 1) Config general
    # -------------------------
    SEED = 42
    N_ENVS = 1  # mantén 1 para debug estable; cuando todo esté perfecto, sube
    # OJO: si usas PyBullet + Subproc a veces se atasca; DummyVecEnv suele ser más estable.

    # -------------------------
    # 2) Hiperparámetros PPO
    # -------------------------
    LR = 1e-4
    GAMMA = 0.99
    N_STEPS = 2048
    BATCH_SIZE = 512
    N_EPOCHS = 10
    ENT_COEF = 0.01
    GAE_LAMBDA = 0.95
    CLIP_RANGE = 0.2

    policy_kwargs = dict(net_arch=dict(pi=[256, 256], vf=[256, 256]))

    # -------------------------
    # 3) Nombre del experimento: timestamp + hparams
    # -------------------------
    timestamp = time.strftime("%Y%m%d-%H%M%S")
    run_name = (
        f"{timestamp}_PPO_curriculum"
        f"_lr{LR:.0e}"
        f"_g{GAMMA}"
        f"_ns{N_STEPS}"
        f"_bs{BATCH_SIZE}"
        f"_ep{N_EPOCHS}"
        f"_ent{ENT_COEF}"
    )

    run_dir = os.path.join("cur_lear", run_name)
    os.makedirs(run_dir, exist_ok=True)

    tb_dir = os.path.join(run_dir, "tb")
    os.makedirs(tb_dir, exist_ok=True)

    # -------------------------
    # 4) Definir etapas del curriculum
    # -------------------------
    stages = [
        dict(name="stage1_find",   env_cls=ThiefEnv_V1a, timesteps=20_000),
        dict(name="stage2_take",   env_cls=ThiefEnv_V1b, timesteps=200_000),
        dict(name="stage3_exit",   env_cls=ThiefEnv_V1c, timesteps=200_000),
    ]

    # Ruta donde iremos guardando el modelo “actual” entre etapas
    current_model_path = os.path.join(run_dir, "model_current.zip")

    # Variable para mantener el modelo en memoria
    model = None

    # -------------------------
    # 5) Loop de entrenamiento por etapas
    # -------------------------
    total_so_far = 0

    for i, st in enumerate(stages, start=1):
        stage_name = st["name"]
        env_cls = st["env_cls"]
        stage_timesteps = st["timesteps"]

        print(f"\n==============================")
        print(f"➡️  Etapa {i}/{len(stages)}: {stage_name}")
        print(f"   Env: {env_cls.__name__}")
        print(f"   Timesteps: {stage_timesteps}")
        print(f"==============================\n")

        # --- Crear env de train y eval para ESTA etapa ---
        vec_env = DummyVecEnv([make_env(env_cls, 0, seed=SEED, render=False)])
        vec_env = VecMonitor(vec_env, filename=os.path.join(run_dir, f"{stage_name}_monitor.csv"))

        eval_env = DummyVecEnv([make_env(env_cls, 0, seed=SEED + 10_000, render=False)])
        eval_env = VecMonitor(eval_env)

        # --- Callbacks específicos de esta etapa ---
        eval_cb = EvalCallback(
            eval_env,
            best_model_save_path=os.path.join(run_dir, f"best_{stage_name}"),
            log_path=os.path.join(run_dir, f"eval_{stage_name}"),
            eval_freq=20_000,
            n_eval_episodes=3,
            deterministic=True,
            render=False,
        )

        checkpoint_cb = CheckpointCallback(
            save_freq=20_000,
            save_path=os.path.join(run_dir, "checkpoints"),
            name_prefix=f"ppo_{stage_name}"
        )

        # --- Crear modelo o cargar y continuar ---
        if model is None:
            # Primera etapa: modelo desde cero
            model = PPO(
                policy="MultiInputPolicy",
                env=vec_env,
                learning_rate=LR,
                n_steps=N_STEPS,
                batch_size=BATCH_SIZE,
                n_epochs=N_EPOCHS,
                gamma=GAMMA,
                gae_lambda=GAE_LAMBDA,
                clip_range=CLIP_RANGE,
                ent_coef=ENT_COEF,
                vf_coef=0.5,
                max_grad_norm=0.5,
                policy_kwargs=policy_kwargs,
                tensorboard_log=tb_dir,
                seed=SEED,
                verbose=1,
            )
        else:
            # Etapas siguientes: cargamos pesos y cambiamos env
            # (PPO.load no “hereda” el env si no se lo pasas)
            model = PPO.load(current_model_path, env=vec_env)

        # --- Entrenar esta etapa ---
        # reset_num_timesteps=False es IMPORTANTE para:
        # - que TensorBoard no “reinicie” el eje x
        # - que el contador global sea continuo entre etapas
        model.learn(
            total_timesteps=stage_timesteps,
            callback=[checkpoint_cb],
            progress_bar=False,
            reset_num_timesteps=False,
            tb_log_name=stage_name,  # separa dashboards por etapa
        )
         
        total_so_far += stage_timesteps

        # --- Guardar “modelo actual” al terminar la etapa ---
        model.save(current_model_path)
        print(f"\n✅ Fin de {stage_name}. Guardado: {current_model_path}")
        print(f"✅ Total acumulado: {total_so_far} timesteps\n")
        
        #Guardar vídeo
        video_dir = os.path.join(run_dir, "videos", stage_name, str(int(time.time())))
        record_one_episode(model, env_cls, video_dir, stage_name=stage_name, seed=SEED)
        print(f"🎥 Video guardado en: {video_dir}")
        '''
        # --- Leer último evaluation (creado por EvalCallback) ---
        eval_dir = os.path.join(run_dir, f"eval_{stage_name}")
        mean_r, std_r = get_last_eval_mean_std(eval_dir)

        # --- Excel dentro de la carpeta juego ---
        base_dir = os.path.dirname(os.path.abspath(__file__))  # carpeta donde está este .py (juego/)
        results_xlsx = os.path.join(base_dir, "results.xlsx")

        append_row_to_excel(
            xlsx_path=results_xlsx,
            sheet_name="curriculum_runs",
            row_dict={
                "run_name": run_name,
                "stage": stage_name,
                "timesteps_stage": stage_timesteps,
                "mean_reward": mean_r,
                "std_reward": std_r,
                "LR": LR,
                "GAMMA": GAMMA,
                "N_STEPS": N_STEPS,
                "BATCH_SIZE": BATCH_SIZE,
                "N_EPOCHS": N_EPOCHS,
                "ENT_COEF": ENT_COEF,
                "GAE_LAMBDA": GAE_LAMBDA,
                "CLIP_RANGE": CLIP_RANGE,
            }
        )

        print(f"📄 Excel actualizado: {results_xlsx} | {stage_name} mean={mean_r} std={std_r}")
        '''

        vec_env.close()
        eval_env.close()

    # -------------------------
    # 6) Guardar modelo final “bonito”
    # -------------------------
    final_path = os.path.join(run_dir, "final_curriculum_model.zip")
    model.save(final_path)
    print(f"\n🏁 Curriculum terminado. Modelo final: {final_path}")

    '''
    # --- Sacar mean/std del último evaluation y guardarlo a Excel ---
    mean_r, std_r = evaluate_policy(
        model,
        eval_env,
        n_eval_episodes=5,
        deterministic=True,
        return_episode_rewards=False
    )

    results_xlsx = os.path.join("curr", "results.xlsx")  # o donde quieras
    append_row_to_excel(
        xlsx_path=results_xlsx,
        sheet_name="curriculum_runs",
        row_dict={
            "run_name": run_name,
            "stage": stage_name,
            "timesteps_stage": stage_timesteps,
            "mean_reward": mean_r,
            "std_reward": std_r,
            "LR": LR,
            "GAMMA": GAMMA,
            "N_STEPS": N_STEPS,
            "BATCH_SIZE": BATCH_SIZE,
            "N_EPOCHS": N_EPOCHS,
            "ENT_COEF": ENT_COEF,
            "GAE_LAMBDA": GAE_LAMBDA,
            "CLIP_RANGE": CLIP_RANGE,
        }
    )
    print(f"📄 Log guardado en Excel: {results_xlsx} (stage={stage_name})")
    '''

if __name__ == "__main__":
    main()
